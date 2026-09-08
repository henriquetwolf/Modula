"""Quarta planilha de teste: 25 nomes comuns com exatamente 1 sobrenome.

Sem regra de proporcao por sexo (nao foi pedida). Os CPFs sao validos e nao
repetem os das planilhas ja geradas, para todas conviverem no mesmo tenant.
"""

from __future__ import annotations

import random
import re
import sys
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

RAIZ = Path(__file__).resolve().parents[1]
ARQUIVO_SAIDA = RAIZ / "lista_teste_comuns_25.xlsx"
PLANILHAS_ANTERIORES = (
    RAIZ / "lista_teste_cpf.xlsx",
    RAIZ / "lista_teste_nomes_incomuns.xlsx",
    RAIZ / "lista_teste_nomes_moderados.xlsx",
)

TOTAL = 25
SEED = 20260922

PRIMEIROS_NOMES = [
    "Ana", "Bruno", "Carla", "Daniel", "Eduardo", "Fernanda", "Gabriel",
    "Isabela", "João", "Juliana", "Larissa", "Lucas", "Marcela", "Mateus",
    "Paula", "Pedro", "Rafael", "Roberta", "Rodrigo", "Sofia", "Thiago",
    "Vanessa", "Vitor", "Camila", "Felipe", "Beatriz", "Gustavo", "Letícia",
    "Marcelo", "Patrícia",
]

SOBRENOMES = [
    "Silva", "Santos", "Oliveira", "Souza", "Lima", "Pereira", "Ferreira",
    "Costa", "Rodrigues", "Almeida", "Nascimento", "Carvalho", "Araújo",
    "Ribeiro", "Gomes", "Martins", "Rocha", "Barbosa", "Mendes", "Cardoso",
    "Teixeira", "Moreira", "Correia", "Dias", "Nunes",
]


def calcular_digito(digitos: list[int]) -> int:
    """Calcula um digito verificador de CPF pelo modulo 11."""
    peso_inicial = len(digitos) + 1
    soma = sum(digito * (peso_inicial - indice) for indice, digito in enumerate(digitos))
    resto = soma % 11
    return 0 if resto < 2 else 11 - resto


def gerar_cpf(rng: random.Random) -> str:
    """Gera um CPF valido de 11 digitos, descartando sequencias repetidas."""
    while True:
        base = [rng.randint(0, 9) for _ in range(9)]
        if len(set(base)) == 1:
            continue
        base.append(calcular_digito(base))
        base.append(calcular_digito(base))
        if len(set(base)) == 1:
            continue
        return "".join(str(digito) for digito in base)


def formatar_cpf(cpf: str) -> str:
    return f"{cpf[:3]}.{cpf[3:6]}.{cpf[6:9]}-{cpf[9:]}"


def cpfs_ja_usados() -> set[str]:
    """CPFs (so digitos) das planilhas anteriores, para nao repetir."""
    usados: set[str] = set()
    for caminho in PLANILHAS_ANTERIORES:
        if not caminho.exists():
            continue
        planilha = load_workbook(caminho, read_only=True).active
        for _, cpf in planilha.iter_rows(min_row=2, max_col=2, values_only=True):
            if cpf:
                usados.add(re.sub(r"\D", "", str(cpf)))
    return usados


def main() -> None:
    rng = random.Random(SEED)

    primeiros = PRIMEIROS_NOMES.copy()
    rng.shuffle(primeiros)
    if TOTAL > len(primeiros):
        raise SystemExit("Faltam primeiros nomes para gerar a lista.")

    linhas: list[tuple[str, str]] = []
    nomes_vistos: set[str] = set()
    cpfs_usados = cpfs_ja_usados()

    for primeiro in primeiros[:TOTAL]:
        # Garante nome completo unico dentro da lista
        while True:
            nome = f"{primeiro} {rng.choice(SOBRENOMES)}"
            if nome not in nomes_vistos:
                break
        nomes_vistos.add(nome)

        while True:
            cpf = gerar_cpf(rng)
            if cpf not in cpfs_usados:
                break
        cpfs_usados.add(cpf)

        linhas.append((nome, formatar_cpf(cpf)))

    workbook = Workbook()
    planilha = workbook.active
    planilha.title = "Nomes comuns"

    planilha.append(["Nome", "CPF"])
    for celula in planilha[1]:
        celula.font = Font(bold=True)
        celula.alignment = Alignment(horizontal="left")

    for nome, cpf in linhas:
        planilha.append([nome, cpf])

    # CPF como texto para nao perder zeros a esquerda em reimportacoes
    for linha in planilha.iter_rows(min_row=2, min_col=2, max_col=2):
        for celula in linha:
            celula.number_format = "@"

    for indice, largura in enumerate((34, 20), start=1):
        planilha.column_dimensions[get_column_letter(indice)].width = largura

    planilha.freeze_panes = "A2"
    workbook.save(ARQUIVO_SAIDA)

    print(f"Arquivo gerado: {ARQUIVO_SAIDA.name}")
    print(f"Total: {len(linhas)} linhas (todos com 1 sobrenome)")


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    main()
