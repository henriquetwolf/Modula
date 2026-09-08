"""Revalida lista_teste_comuns_25.xlsx a partir do arquivo salvo em disco.

Confere CPFs (digitos verificadores), unicidade, 1 sobrenome por nome e ausencia
de colisao de CPF com as planilhas anteriores.
"""

from __future__ import annotations

import re
import sys
from pathlib import Path

from openpyxl import load_workbook

RAIZ = Path(__file__).resolve().parents[1]
ARQUIVO = RAIZ / "lista_teste_comuns_25.xlsx"
PLANILHAS_ANTERIORES = (
    RAIZ / "lista_teste_cpf.xlsx",
    RAIZ / "lista_teste_nomes_incomuns.xlsx",
    RAIZ / "lista_teste_nomes_moderados.xlsx",
)

PADRAO_CPF = re.compile(r"^\d{3}\.\d{3}\.\d{3}-\d{2}$")


def cpf_valido(cpf_formatado: str) -> bool:
    if not PADRAO_CPF.match(cpf_formatado):
        return False
    digitos = [int(c) for c in re.sub(r"\D", "", cpf_formatado)]
    if len(set(digitos)) == 1:
        return False
    for tamanho in (9, 10):
        peso = tamanho + 1
        soma = sum(d * (peso - i) for i, d in enumerate(digitos[:tamanho]))
        resto = soma % 11
        esperado = 0 if resto < 2 else 11 - resto
        if digitos[tamanho] != esperado:
            return False
    return True


def main() -> None:
    planilha = load_workbook(ARQUIVO, read_only=True).active
    linhas = list(planilha.iter_rows(min_row=1, max_col=2, values_only=True))

    cabecalho, dados = linhas[0], linhas[1:]
    assert cabecalho == ("Nome", "CPF"), f"Cabecalho inesperado: {cabecalho}"
    assert len(dados) == 25, f"Esperado 25 linhas, obtido {len(dados)}"

    nomes = [linha[0] for linha in dados]
    cpfs = [linha[1] for linha in dados]

    assert len(set(nomes)) == 25, "Nomes duplicados"
    assert len(set(cpfs)) == 25, "CPFs duplicados"

    # Todo nome deve ter exatamente 2 palavras: 1 primeiro nome + 1 sobrenome
    for nome in nomes:
        assert len(nome.split()) == 2, f"Nome sem exatamente 1 sobrenome: {nome}"

    invalidos = [cpf for cpf in cpfs if not cpf_valido(cpf)]
    assert not invalidos, f"CPFs invalidos: {invalidos}"

    usados: set[str] = set()
    for caminho in PLANILHAS_ANTERIORES:
        if not caminho.exists():
            continue
        anterior = load_workbook(caminho, read_only=True).active
        for _, cpf in anterior.iter_rows(min_row=2, max_col=2, values_only=True):
            if cpf:
                usados.add(re.sub(r"\D", "", str(cpf)))

    colisoes = [cpf for cpf in cpfs if re.sub(r"\D", "", cpf) in usados]
    assert not colisoes, f"CPFs repetidos de outra planilha: {colisoes}"

    print("OK: 25 linhas, todos com 1 sobrenome, CPFs validos, unicos e sem colisao.")


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    main()
