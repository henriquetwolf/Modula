"""Terceira planilha de teste: primeiros nomes de uso corrente (nem raros, nem
os mais batidos) com 15 pessoas de 1 sobrenome, 15 de 2 e 15 de 3, sendo 40%
feminino.

Os CPFs sao validos e nao repetem os das planilhas ja geradas, para as tres
listas poderem conviver no mesmo tenant.
"""

from __future__ import annotations

import random
import re
import sys
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

RAIZ = Path(__file__).resolve().parents[1]
ARQUIVO_SAIDA = RAIZ / "lista_teste_nomes_moderados.xlsx"
PLANILHAS_ANTERIORES = (
    RAIZ / "lista_teste_cpf.xlsx",
    RAIZ / "lista_teste_nomes_incomuns.xlsx",
)

POR_GRUPO = 15
QUANTIDADES_SOBRENOMES = (1, 2, 3)
PROPORCAO_FEMININO = 0.40
SEED = 20260915

# Sobrenomes com particula ("de", "da") ficaram de fora de proposito para a
# contagem de sobrenomes continuar exata.
PRIMEIROS_NOMES_FEMININOS = [
    "Adriana", "Alice", "Bianca", "Carla", "Cecília", "Cristiane", "Débora",
    "Eliane", "Flávia", "Heloísa", "Ingrid", "Jussara", "Kelly", "Lorena",
    "Márcia", "Michele", "Nádia", "Priscila", "Renata", "Rosana", "Silvana",
    "Simone", "Tânia", "Viviane",
]

PRIMEIROS_NOMES_MASCULINOS = [
    "Alexandre", "Anderson", "Cláudio", "Cristiano", "Douglas", "Edson",
    "Emerson", "Everton", "Fabiano", "Fábio", "Gilberto", "Hélio", "Ivan",
    "Jefferson", "Júlio", "Luciano", "Márcio", "Maurício", "Nelson", "Osmar",
    "Osvaldo", "Reinaldo", "Renato", "Ricardo", "Roberto", "Sandro", "Sérgio",
    "Silvio", "Válter", "Wagner", "Wesley", "Wilson",
]

SOBRENOMES = [
    "Almeida", "Alves", "Araújo", "Barbosa", "Barros", "Batista", "Cardoso",
    "Carvalho", "Castro", "Correia", "Costa", "Dias", "Duarte", "Fernandes",
    "Ferreira", "Freitas", "Gomes", "Gonçalves", "Lima", "Lopes", "Machado",
    "Martins", "Melo", "Mendes", "Monteiro", "Moraes", "Moreira",
    "Nascimento", "Nunes", "Oliveira", "Pereira", "Pinto", "Ramos", "Ribeiro",
    "Rocha", "Rodrigues", "Santos", "Silva", "Souza", "Teixeira", "Vieira",
]


def calcular_digito(digitos: list[int]) -> int:
    """Calcula um digito verificador de CPF pelo modulo 11."""
    peso_inicial = len(digitos) + 1
    soma = sum(digito * (peso_inicial - indice) for indice, digito in enumerate(digitos))
    resto = soma % 11
    return 0 if resto < 2 else 11 - resto


def gerar_cpf(rng: random.Random) -> str:
    """Gera um CPF valido de 11 digitos, descartando sequencias repetidas."""
    while True:
        base = [rng.randint(0, 9) for _ in range(9)]
        if len(set(base)) == 1:
            continue
        base.append(calcular_digito(base))
        base.append(calcular_digito(base))
        if len(set(base)) == 1:
            continue
        return "".join(str(digito) for digito in base)


def formatar_cpf(cpf: str) -> str:
    return f"{cpf[:3]}.{cpf[3:6]}.{cpf[6:9]}-{cpf[9:]}"


def cpfs_ja_usados() -> set[str]:
    """CPFs (so digitos) das planilhas anteriores, para nao repetir."""
    usados: set[str] = set()
    for caminho in PLANILHAS_ANTERIORES:
        if not caminho.exists():
            continue
        planilha = load_workbook(caminho, read_only=True).active
        for _, cpf in planilha.iter_rows(min_row=2, max_col=2, values_only=True):
            if cpf:
                usados.add(re.sub(r"\D", "", str(cpf)))
    return usados


def main() -> None:
    rng = random.Random(SEED)

    femininos = PRIMEIROS_NOMES_FEMININOS.copy()
    masculinos = PRIMEIROS_NOMES_MASCULINOS.copy()
    rng.shuffle(femininos)
    rng.shuffle(masculinos)

    qtd_feminino = round(POR_GRUPO * PROPORCAO_FEMININO)
    qtd_masculino = POR_GRUPO - qtd_feminino

    if qtd_feminino * len(QUANTIDADES_SOBRENOMES) > len(femininos):
        raise SystemExit("Faltam primeiros nomes femininos.")
    if qtd_masculino * len(QUANTIDADES_SOBRENOMES) > len(masculinos):
        raise SystemExit("Faltam primeiros nomes masculinos.")

    linhas: list[tuple[str, str, int, str]] = []
    cpfs_usados = cpfs_ja_usados()

    for total_sobrenomes in QUANTIDADES_SOBRENOMES:
        for sexo, quantidade in (("F", qtd_feminino), ("M", qtd_masculino)):
            reserva = femininos if sexo == "F" else masculinos
            for _ in range(quantidade):
                primeiro = reserva.pop()
                sobrenomes = rng.sample(SOBRENOMES, total_sobrenomes)
                nome = " ".join([primeiro, *sobrenomes])

                while True:
                    cpf = gerar_cpf(rng)
                    if cpf not in cpfs_usados:
                        break
                cpfs_usados.add(cpf)

                linhas.append((nome, formatar_cpf(cpf), total_sobrenomes, sexo))

    # Embaralha para os grupos e os sexos nao sairem agrupados na planilha
    rng.shuffle(linhas)

    workbook = Workbook()
    planilha = workbook.active
    planilha.title = "Nomes correntes"

    planilha.append(["Nome", "CPF"])
    for celula in planilha[1]:
        celula.font = Font(bold=True)
        celula.alignment = Alignment(horizontal="left")

    for nome, cpf, _, _ in linhas:
        planilha.append([nome, cpf])

    # CPF como texto para nao perder zeros a esquerda em reimportacoes
    for linha in planilha.iter_rows(min_row=2, min_col=2, max_col=2):
        for celula in linha:
            celula.number_format = "@"

    for indice, largura in enumerate((46, 20), start=1):
        planilha.column_dimensions[get_column_letter(indice)].width = largura

    planilha.freeze_panes = "A2"
    workbook.save(ARQUIVO_SAIDA)

    total_f = sum(1 for _, _, _, sexo in linhas if sexo == "F")
    print(f"Arquivo gerado: {ARQUIVO_SAIDA.name}")
    print(f"Total: {len(linhas)} linhas")
    for total_sobrenomes in QUANTIDADES_SOBRENOMES:
        grupo = [l for l in linhas if l[2] == total_sobrenomes]
        f = sum(1 for l in grupo if l[3] == "F")
        print(f"  {total_sobrenomes} sobrenome(s): {len(grupo)} pessoas ({f} F / {len(grupo) - f} M)")
    print(f"Feminino: {total_f} de {len(linhas)} ({total_f / len(linhas):.0%})")


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    main()
