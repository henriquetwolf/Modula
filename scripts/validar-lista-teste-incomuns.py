"""Revalida lista_teste_nomes_incomuns.xlsx a partir do arquivo salvo em disco.

Confere CPFs (digitos verificadores), unicidade, ausencia de colisao com a
primeira planilha, a contagem de sobrenomes por grupo e a proporcao feminina.
"""

from __future__ import annotations

import re
import sys
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook

RAIZ = Path(__file__).resolve().parents[1]
ARQUIVO = RAIZ / "lista_teste_nomes_incomuns.xlsx"
ARQUIVO_ANTERIOR = RAIZ / "lista_teste_cpf.xlsx"

PADRAO_CPF = re.compile(r"^\d{3}\.\d{3}\.\d{3}-\d{2}$")

# Repetido de proposito: a validacao nao deve depender do gerador.
FEMININOS = {
    "Adelvina", "Belmira", "Corina", "Deolinda", "Ercília", "Filomena",
    "Genoveva", "Hermínia", "Idalina", "Jacinta", "Leocádia", "Marcolina",
    "Nazaré", "Olegária", "Placidina", "Quitéria", "Romilda", "Saturnina",
    "Teodolinda", "Ubaldina", "Valquíria", "Wanderleia", "Xênia", "Zulmira",
}


def cpf_valido(cpf_formatado: str) -> bool:
    if not PADRAO_CPF.match(cpf_formatado):
        return False
    digitos = [int(c) for c in re.sub(r"\D", "", cpf_formatado)]
    if len(set(digitos)) == 1:
        return False
    for tamanho in (9, 10):
        peso = tamanho + 1
        soma = sum(d * (peso - i) for i, d in enumerate(digitos[:tamanho]))
        resto = soma % 11
        esperado = 0 if resto < 2 else 11 - resto
        if digitos[tamanho] != esperado:
            return False
    return True


def main() -> None:
    planilha = load_workbook(ARQUIVO, read_only=True).active
    linhas = list(planilha.iter_rows(min_row=1, max_col=2, values_only=True))

    cabecalho, dados = linhas[0], linhas[1:]
    assert cabecalho == ("Nome", "CPF"), f"Cabecalho inesperado: {cabecalho}"
    assert len(dados) == 45, f"Esperado 45 linhas, obtido {len(dados)}"

    nomes = [linha[0] for linha in dados]
    cpfs = [linha[1] for linha in dados]

    assert len(set(nomes)) == 45, "Nomes duplicados"
    assert len(set(cpfs)) == 45, "CPFs duplicados"

    invalidos = [cpf for cpf in cpfs if not cpf_valido(cpf)]
    assert not invalidos, f"CPFs invalidos: {invalidos}"

    # Nenhum CPF pode coincidir com a primeira planilha
    if ARQUIVO_ANTERIOR.exists():
        anterior = load_workbook(ARQUIVO_ANTERIOR, read_only=True).active
        usados = {
            re.sub(r"\D", "", str(cpf))
            for _, cpf in anterior.iter_rows(min_row=2, max_col=2, values_only=True)
            if cpf
        }
        colisoes = [cpf for cpf in cpfs if re.sub(r"\D", "", cpf) in usados]
        assert not colisoes, f"CPFs repetidos da outra planilha: {colisoes}"

    # Contagem de sobrenomes: 1 sobrenome = 2 palavras, 2 = 3 palavras, 3 = 4
    por_qtd = Counter(len(nome.split()) - 1 for nome in nomes)
    assert por_qtd == {1: 15, 2: 15, 3: 15}, f"Distribuicao de sobrenomes errada: {dict(por_qtd)}"

    # Nenhum sobrenome repetido dentro do mesmo nome
    for nome in nomes:
        sobrenomes = nome.split()[1:]
        assert len(set(sobrenomes)) == len(sobrenomes), f"Sobrenome repetido em: {nome}"

    femininos = [nome for nome in nomes if nome.split()[0] in FEMININOS]
    assert len(femininos) == 18, f"Esperado 18 femininos, obtido {len(femininos)}"

    print("OK: 45 linhas, CPFs validos e unicos, sem colisao com a outra planilha.")
    print(f"Sobrenomes: 1 -> {por_qtd[1]}, 2 -> {por_qtd[2]}, 3 -> {por_qtd[3]}")
    print(f"Feminino: {len(femininos)}/45 ({len(femininos) / 45:.0%})")
    print("\nAmostra por grupo:")
    for qtd in (1, 2, 3):
        exemplo = next(n for n in nomes if len(n.split()) - 1 == qtd)
        print(f"  {qtd} sobrenome(s): {exemplo}")


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    main()
