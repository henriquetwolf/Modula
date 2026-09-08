"""Revalida a planilha lista_teste_cpf.xlsx: unicidade e digitos verificadores."""

from __future__ import annotations

import re
from pathlib import Path

from openpyxl import load_workbook

ARQUIVO = Path(__file__).resolve().parents[1] / "lista_teste_cpf.xlsx"
PADRAO = re.compile(r"^\d{3}\.\d{3}\.\d{3}-\d{2}$")


def cpf_valido(cpf_formatado: str) -> bool:
    if not PADRAO.match(cpf_formatado):
        return False
    digitos = [int(caractere) for caractere in re.sub(r"\D", "", cpf_formatado)]
    if len(set(digitos)) == 1:
        return False
    for tamanho in (9, 10):
        peso_inicial = tamanho + 1
        soma = sum(digito * (peso_inicial - i) for i, digito in enumerate(digitos[:tamanho]))
        resto = soma % 11
        esperado = 0 if resto < 2 else 11 - resto
        if digitos[tamanho] != esperado:
            return False
    return True


def main() -> None:
    planilha = load_workbook(ARQUIVO, read_only=True).active
    linhas = list(planilha.iter_rows(min_row=1, max_col=2, values_only=True))

    cabecalho, dados = linhas[0], linhas[1:]
    assert cabecalho == ("Nome", "CPF"), f"Cabecalho inesperado: {cabecalho}"
    assert len(dados) == 50, f"Esperado 50 linhas de dados, obtido {len(dados)}"

    nomes = [linha[0] for linha in dados]
    cpfs = [linha[1] for linha in dados]

    assert all(nome and len(nome.split()) >= 3 for nome in nomes), "Nome incompleto encontrado"
    assert len(set(nomes)) == 50, "Nomes duplicados encontrados"
    assert len(set(cpfs)) == 50, "CPFs duplicados encontrados"

    invalidos = [cpf for cpf in cpfs if not cpf_valido(cpf)]
    assert not invalidos, f"CPFs invalidos: {invalidos}"

    print("OK: 50 nomes completos unicos e 50 CPFs validos unicos.")
    print("Amostra:")
    for nome, cpf in dados[:5]:
        print(f"  {nome:<32} {cpf}")


if __name__ == "__main__":
    main()
