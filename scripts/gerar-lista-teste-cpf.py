"""Gera uma planilha .xlsx com nomes completos ficticios e CPFs validos para testes."""

from __future__ import annotations

import random
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

TOTAL_LINHAS = 50
PROPORCAO_MULHERES = 0.80
SEED = 20260908
ARQUIVO_SAIDA = Path(__file__).resolve().parents[1] / "lista_teste_cpf.xlsx"

PRIMEIROS_NOMES_FEMININOS = [
    "Ana", "Beatriz", "Camila", "Daniela", "Eduarda", "Fernanda", "Gabriela",
    "Helena", "Isabela", "Juliana", "Larissa", "Mariana", "Natalia", "Patricia",
    "Rafaela", "Sabrina", "Tatiane", "Vanessa", "Amanda", "Bruna", "Carolina",
]

PRIMEIROS_NOMES_MASCULINOS = [
    "Andre", "Bruno", "Caio", "Daniel", "Eduardo", "Felipe", "Gustavo",
    "Henrique", "Igor", "Joao", "Leonardo", "Marcelo", "Nicolas", "Otavio",
    "Pedro", "Rodrigo", "Samuel", "Thiago", "Vinicius", "Lucas", "Matheus",
]

NOMES_MEIO = [
    "Alves", "Barbosa", "Cardoso", "Duarte", "Ferreira", "Gomes", "Henrique",
    "Isabel", "Jose", "Lima", "Maria", "Nunes", "Oliveira", "Pereira",
    "Queiroz", "Ribeiro", "Santos", "Teixeira", "Vieira", "Moraes",
]

SOBRENOMES = [
    "Almeida", "Azevedo", "Batista", "Cavalcanti", "Correia", "Costa", "Dias",
    "Fonseca", "Freitas", "Machado", "Martins", "Medeiros", "Mendes", "Monteiro",
    "Pinheiro", "Rocha", "Rodrigues", "Sales", "Siqueira", "Souza", "Tavares",
    "Carvalho", "Nascimento", "Bezerra", "Andrade",
]


def calcular_digito(digitos: list[int]) -> int:
    """Calcula um digito verificador de CPF pelo modulo 11."""
    peso_inicial = len(digitos) + 1
    soma = sum(digito * (peso_inicial - indice) for indice, digito in enumerate(digitos))
    resto = soma % 11
    return 0 if resto < 2 else 11 - resto


def gerar_cpf(rng: random.Random) -> str:
    """Gera um CPF valido de 11 digitos, descartando sequencias repetidas."""
    while True:
        base = [rng.randint(0, 9) for _ in range(9)]
        if len(set(base)) == 1:
            continue
        base.append(calcular_digito(base))
        base.append(calcular_digito(base))
        if len(set(base)) == 1:
            continue
        return "".join(str(digito) for digito in base)


def formatar_cpf(cpf: str) -> str:
    return f"{cpf[:3]}.{cpf[3:6]}.{cpf[6:9]}-{cpf[9:]}"


def gerar_nome(rng: random.Random, primeiros: list[str]) -> str:
    return " ".join(
        [rng.choice(primeiros), rng.choice(NOMES_MEIO), rng.choice(SOBRENOMES)]
    )


def main() -> None:
    rng = random.Random(SEED)

    qtd_mulheres = round(TOTAL_LINHAS * PROPORCAO_MULHERES)
    qtd_homens = TOTAL_LINHAS - qtd_mulheres

    nomes: list[str] = []
    vistos_nomes: set[str] = set()
    for primeiros, quantidade in (
        (PRIMEIROS_NOMES_FEMININOS, qtd_mulheres),
        (PRIMEIROS_NOMES_MASCULINOS, qtd_homens),
    ):
        gerados = 0
        while gerados < quantidade:
            nome = gerar_nome(rng, primeiros)
            if nome in vistos_nomes:
                continue
            vistos_nomes.add(nome)
            nomes.append(nome)
            gerados += 1

    rng.shuffle(nomes)

    cpfs: list[str] = []
    vistos_cpfs: set[str] = set()
    while len(cpfs) < TOTAL_LINHAS:
        cpf = gerar_cpf(rng)
        if cpf in vistos_cpfs:
            continue
        vistos_cpfs.add(cpf)
        cpfs.append(formatar_cpf(cpf))

    workbook = Workbook()
    planilha = workbook.active
    planilha.title = "Lista de teste"

    planilha.append(["Nome", "CPF"])
    for celula in planilha[1]:
        celula.font = Font(bold=True)
        celula.alignment = Alignment(horizontal="left")

    for nome, cpf in zip(nomes, cpfs):
        planilha.append([nome, cpf])

    # CPF fica como texto para nao perder zeros a esquerda em reimportacoes.
    for linha in planilha.iter_rows(min_row=2, min_col=2, max_col=2):
        for celula in linha:
            celula.number_format = "@"

    for indice, largura in enumerate((34, 20), start=1):
        planilha.column_dimensions[get_column_letter(indice)].width = largura

    planilha.freeze_panes = "A2"

    workbook.save(ARQUIVO_SAIDA)
    print(f"Arquivo gerado: {ARQUIVO_SAIDA}")
    print(f"Linhas de dados: {planilha.max_row - 1}")


if __name__ == "__main__":
    main()
