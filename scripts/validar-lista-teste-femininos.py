"""Revalida lista_teste_femininos.xlsx a partir do arquivo salvo em disco.

Confere CPFs (digitos verificadores), unicidade, a contagem por sobrenomes
(25 com 1 e 25 com 3), que todos os primeiros nomes sejam femininos e a
ausencia de colisao de CPF com as planilhas anteriores.
"""

from __future__ import annotations

import re
import sys
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook

RAIZ = Path(__file__).resolve().parents[1]
ARQUIVO = RAIZ / "lista_teste_femininos.xlsx"
PLANILHAS_ANTERIORES = (
    RAIZ / "lista_teste_cpf.xlsx",
    RAIZ / "lista_teste_nomes_incomuns.xlsx",
    RAIZ / "lista_teste_nomes_moderados.xlsx",
    RAIZ / "lista_teste_comuns_25.xlsx",
)

PADRAO_CPF = re.compile(r"^\d{3}\.\d{3}\.\d{3}-\d{2}$")

# Repetido de proposito: a validacao nao deve depender do gerador.
FEMININOS = {
    "Ana", "Beatriz", "Camila", "Carla", "Carolina", "Cristiane", "Daniela",
    "Débora", "Eduarda", "Fernanda", "Gabriela", "Helena", "Isabela",
    "Jéssica", "Juliana", "Larissa", "Letícia", "Luana", "Lúcia", "Marcela",
    "Mariana", "Michele", "Natália", "Patrícia", "Paula", "Priscila",
    "Rafaela", "Renata", "Roberta", "Sabrina", "Simone", "Tatiane", "Vanessa",
    "Viviane", "Amanda", "Bruna", "Aline", "Adriana", "Bianca",
}


def cpf_valido(cpf_formatado: str) -> bool:
    if not PADRAO_CPF.match(cpf_formatado):
        return False
    digitos = [int(c) for c in re.sub(r"\D", "", cpf_formatado)]
    if len(set(digitos)) == 1:
        return False
    for tamanho in (9, 10):
        peso = tamanho + 1
        soma = sum(d * (peso - i) for i, d in enumerate(digitos[:tamanho]))
        resto = soma % 11
        esperado = 0 if resto < 2 else 11 - resto
        if digitos[tamanho] != esperado:
            return False
    return True


def main() -> None:
    planilha = load_workbook(ARQUIVO, read_only=True).active
    linhas = list(planilha.iter_rows(min_row=1, max_col=2, values_only=True))

    cabecalho, dados = linhas[0], linhas[1:]
    assert cabecalho == ("Nome", "CPF"), f"Cabecalho inesperado: {cabecalho}"
    assert len(dados) == 50, f"Esperado 50 linhas, obtido {len(dados)}"

    nomes = [linha[0] for linha in dados]
    cpfs = [linha[1] for linha in dados]

    assert len(set(nomes)) == 50, "Nomes duplicados"
    assert len(set(cpfs)) == 50, "CPFs duplicados"

    invalidos = [cpf for cpf in cpfs if not cpf_valido(cpf)]
    assert not invalidos, f"CPFs invalidos: {invalidos}"

    # Contagem por sobrenomes: 1 sobrenome = 2 palavras, 3 sobrenomes = 4 palavras
    por_qtd = Counter(len(nome.split()) - 1 for nome in nomes)
    assert por_qtd == {1: 25, 3: 25}, f"Distribuicao de sobrenomes errada: {dict(por_qtd)}"

    # Todos os primeiros nomes devem ser femininos
    nao_femininos = sorted({nome.split()[0] for nome in nomes if nome.split()[0] not in FEMININOS})
    assert not nao_femininos, f"Primeiros nomes nao reconhecidos como femininos: {nao_femininos}"

    # Nenhum sobrenome repetido dentro do mesmo nome
    for nome in nomes:
        sobrenomes = nome.split()[1:]
        assert len(set(sobrenomes)) == len(sobrenomes), f"Sobrenome repetido em: {nome}"

    usados: set[str] = set()
    for caminho in PLANILHAS_ANTERIORES:
        if not caminho.exists():
            continue
        anterior = load_workbook(caminho, read_only=True).active
        for _, cpf in anterior.iter_rows(min_row=2, max_col=2, values_only=True):
            if cpf:
                usados.add(re.sub(r"\D", "", str(cpf)))

    colisoes = [cpf for cpf in cpfs if re.sub(r"\D", "", cpf) in usados]
    assert not colisoes, f"CPFs repetidos de outra planilha: {colisoes}"

    print("OK: 50 linhas femininas, CPFs validos e unicos, sem colisao.")
    print(f"Sobrenomes: 1 -> {por_qtd[1]}, 3 -> {por_qtd[3]}")


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    main()
